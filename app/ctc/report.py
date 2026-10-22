"""CTC report rows and CSV/XLSX export."""
from __future__ import annotations

import csv
import io
from collections import defaultdict
from dataclasses import asdict, dataclass
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.ctc.engine import compute_for_subcontract
from app.db.models import Project, Subcontract

ZERO = Decimal("0")


@dataclass
class ReportRow:
    project_number: str
    project_name: str
    pm: str
    vendor: str
    subcontract_id: int
    original_contract: Decimal
    approved_cos: Decimal
    pending_cos: Decimal
    revised_contract: Decimal
    billed: Decimal
    paid: Decimal
    hold: Decimal
    retainage_held: Decimal
    ctc: Decimal
    pct_complete: Decimal | None
    status: str
    open_waivers: int = 0
    last_waiver_status: str = ""


MONEY_FIELDS = ("original_contract", "approved_cos", "pending_cos", "revised_contract", "billed", "paid", "hold",
                "retainage_held", "ctc")

HEADERS = {
    "project_number": "Project #", "project_name": "Project", "pm": "PM", "vendor": "Sub",
    "original_contract": "Original contract", "approved_cos": "Approved COs", "pending_cos": "Pending COs",
    "revised_contract": "Revised contract", "billed": "Billed", "paid": "Paid", "hold": "On hold",
    "retainage_held": "Retainage held", "ctc": "Cost to complete", "pct_complete": "% complete",
    "status": "Status", "open_waivers": "Open waivers", "last_waiver_status": "Last waiver received",
}


def report_rows(db: Session, project_number: str | None = None) -> list[ReportRow]:
    q = (select(Subcontract).join(Project)
         .where(Project.is_tracked.is_(True), Subcontract.deleted_in_source.is_(False))
         .options(selectinload(Subcontract.payables), selectinload(Subcontract.change_orders),
                  selectinload(Subcontract.vendor), selectinload(Subcontract.project))
         .order_by(Project.number))
    if project_number:
        q = q.where(Project.number == project_number)
    from app.waivers.service import sub_waiver_summary

    summary = sub_waiver_summary(db)
    rows = []
    for sc in db.scalars(q):
        c = compute_for_subcontract(sc)
        rows.append(ReportRow(
            project_number=sc.project.number, project_name=sc.project.name or "",
            pm=sc.project.pm_name or "", vendor=sc.vendor.name, subcontract_id=sc.id,
            original_contract=c.original_contract, approved_cos=c.approved_cos, pending_cos=c.pending_cos,
            revised_contract=c.revised_contract, billed=c.billed_to_date, paid=c.paid_to_date,
            hold=c.hold_total, retainage_held=max(c.retention_owed or ZERO, ZERO), ctc=c.cost_to_complete,
            pct_complete=c.pct_complete, status=c.status,
            open_waivers=summary.get(sc.id, {}).get("open", 0),
            last_waiver_status=summary.get(sc.id, {}).get("last_received", ""),
        ))
    rows.sort(key=lambda r: (r.project_number, r.vendor.lower()))
    return rows


def grouped(rows: list[ReportRow], by: str) -> list[dict]:
    """Totals by job, sub or PM."""
    key = {"job": lambda r: (r.project_number, r.project_name), "sub": lambda r: (r.vendor, ""),
           "pm": lambda r: (r.pm or "(no PM)", "")}[by]
    acc: dict[tuple, dict] = defaultdict(lambda: {f: ZERO for f in MONEY_FIELDS} | {"count": 0, "over_billed": 0})
    for r in rows:
        g = acc[key(r)]
        for f in MONEY_FIELDS:
            g[f] += getattr(r, f)
        g["count"] += 1
        g["over_billed"] += r.status == "over_billed"
    out = []
    for (k1, k2), g in sorted(acc.items()):
        g["key"], g["label"] = k1, k2
        g["pct_complete"] = (g["billed"] / g["revised_contract"]).quantize(Decimal("0.0001")) \
            if g["revised_contract"] > 0 else None
        out.append(g)
    return out


def _detail_table(rows: list[ReportRow]) -> tuple[list[str], list[list]]:
    fields = [f for f in HEADERS]
    data = []
    for r in rows:
        d = asdict(r)
        data.append([d[f] for f in fields])
    return [HEADERS[f] for f in fields], data


def to_csv(rows: list[ReportRow]) -> str:
    headers, data = _detail_table(rows)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    for row in data:
        w.writerow(["" if v is None else v for v in row])
    return buf.getvalue()


def to_xlsx(rows: list[ReportRow]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "By sub on job"
    headers, data = _detail_table(rows)
    _write_sheet(ws, headers, data)
    for by, title, first in (("job", "By job", "Project #"), ("sub", "By sub", "Sub"), ("pm", "By PM", "PM")):
        g = grouped(rows, by)
        hdr = [first] + (["Project"] if by == "job" else []) + [HEADERS[f] for f in MONEY_FIELDS] + \
              ["% complete", "Subcontracts", "Over-billed"]
        body = [[x["key"]] + ([x["label"]] if by == "job" else []) + [x[f] for f in MONEY_FIELDS] +
                [x["pct_complete"], x["count"], x["over_billed"]] for x in g]
        _write_sheet(wb.create_sheet(title), hdr, body)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_sheet(ws, headers, data):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in data:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    for i, h in enumerate(headers, start=1):
        if h == "% complete":
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in cell:
                    c.number_format = "0.0%"
        elif h in {HEADERS[f] for f in MONEY_FIELDS}:
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in cell:
                    c.number_format = '"$"#,##0.00'
    ws.freeze_panes = "A2"
