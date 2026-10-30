"""A sub's suppliers on a job, their lists, and their waivers (which always come through the sub).

The app never emails a supplier. It asks major-trade subs for their supplier list once, reads
the reply with Claude into suggestions a person says yes or no to, asks subs to send their
suppliers' waivers with their own, and files the supplier waivers it recognises.
"""
from __future__ import annotations

import json
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation

import structlog
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db.models import (
    Contact,
    EmailAttachment,
    EmailMessage,
    Project,
    Subcontract,
    Supplier,
    SupplierCandidate,
    Waiver,
)
from app.matching.normalize import normalize_vendor
from app.services.audit import audit
from app.services.reviews import upsert_review
from app.services.settings_store import get_setting
from app.waivers.types import WaiverStatus, WaiverType

log = structlog.get_logger()

SIGNED = (WaiverStatus.RECEIVED.value, WaiverStatus.ACCEPTED.value)


def now() -> datetime:
    return datetime.now(timezone.utc)


def to_amount(v) -> Decimal | None:
    if v in (None, ""):
        return None
    try:
        return Decimal("".join(ch for ch in str(v) if ch.isdigit() or ch == ".")).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


# ---------------------------------------------------------------- the list

def tracked(sc: Subcontract) -> bool:
    """Suppliers can be switched off per job (Jobs page checkbox)."""
    return bool(sc.project.track_suppliers)


def active(db: Session, sc: Subcontract) -> list[Supplier]:
    """The sub's suppliers on this job; none while suppliers are switched off for the job."""
    if not tracked(sc):
        return []
    return list(db.scalars(select(Supplier).where(Supplier.subcontract_id == sc.id, Supplier.status == "active")
                           .order_by(Supplier.name)))


def is_major_trade(db: Session, sc: Subcontract) -> bool:
    """Revised contract at or over the supplier minimum, and not labor only."""
    from app.ctc.engine import compute_for_subcontract

    if (sc.supplier_contractor or "").strip().lower() == "labor" or sc.deleted_in_source or not tracked(sc):
        return False
    return compute_for_subcontract(sc).revised_contract >= Decimal(str(get_setting(db, "supplier_min")))


def _contact_for(db: Session, s: Supplier) -> None:
    """A supplier's own contact goes in the contacts list, under the supplier's company."""
    if not s.email:
        return
    addr = s.email.strip().lower()
    c = db.scalar(select(Contact).where(Contact.email == addr))
    if c is None:
        c = Contact(email=addr, source="supplier", message_count=0, first_seen_at=now())
        db.add(c)
    c.vendor_id = c.vendor_id or s.vendor_id
    c.name = c.name or s.contact_name
    c.phone = c.phone or s.phone
    if c.status != "ignored":
        c.status, c.confirmed_by, c.confirmed_at = "confirmed", c.confirmed_by or "system:supplier-list", \
            c.confirmed_at or now()


def add(db: Session, sc: Subcontract, actor: str, name: str, supplies: str | None = None, amount=None,
        contact_name: str | None = None, email: str | None = None, phone: str | None = None,
        source: str = "manual") -> Supplier:
    from app.matching import resolver

    name = name.strip()
    if not name:
        raise ValueError("Give the supplier's company name.")
    vendor = resolver.get_or_create_vendor(db, name)
    if vendor.type == "sub" and not db.scalar(select(Subcontract.id).where(Subcontract.vendor_id == vendor.id).limit(1)):
        vendor.type = "supplier"
    existing = db.scalar(select(Supplier).where(Supplier.subcontract_id == sc.id, Supplier.vendor_id == vendor.id))
    s = existing or Supplier(subcontract_id=sc.id, vendor_id=vendor.id, name=name, source=source, created_by=actor)
    if existing is None:
        db.add(s)
    s.status, s.removed_at = "active", None
    s.supplies = (supplies or "").strip() or s.supplies
    s.approx_amount = to_amount(amount) if amount not in (None, "") else s.approx_amount
    s.contact_name = (contact_name or "").strip() or s.contact_name
    s.email = (email or "").strip().lower() or s.email
    s.phone = (phone or "").strip() or s.phone
    db.flush()
    _contact_for(db, s)
    if sc.supplier_list_status in (None, "requested", "none"):
        sc.supplier_list_status = "received"
    audit(db, actor, "supplier.added", "subcontract", sc.id, supplier=name, source=source)
    return s


def edit(db: Session, s: Supplier, actor: str, **fields) -> None:
    for k in ("name", "supplies", "contact_name", "phone"):
        if fields.get(k) is not None:
            setattr(s, k, fields[k].strip() or None if k != "name" else (fields[k].strip() or s.name))
    if fields.get("email") is not None:
        s.email = fields["email"].strip().lower() or None
    if fields.get("amount") is not None:
        s.approx_amount = to_amount(fields["amount"])
    _contact_for(db, s)
    audit(db, actor, "supplier.edited", "supplier", s.id)


def remove(db: Session, s: Supplier, actor: str) -> None:
    s.status, s.removed_at = "removed", now()
    audit(db, actor, "supplier.removed", "supplier", s.id, name=s.name)
    recheck_finals(db, s.subcontract)


def mark_no_suppliers(db: Session, sc: Subcontract, actor: str) -> None:
    sc.supplier_list_status = "none"
    audit(db, actor, "supplier.list_none", "subcontract", sc.id)


# ---------------------------------------------------------------- asking for the list

def request_list(db: Session, sc: Subcontract, actor: str = "system:suppliers") -> EmailMessage | None:
    """Draft the supplier-list request (template 09). It goes to the Outbox like other mail."""
    from app.mail import outbound

    e = outbound.queue_supplier_list(db, sc, actor)
    sc.supplier_list_status = "requested"
    sc.supplier_list_requested_at = now()
    sc.supplier_list_reminders = 0
    audit(db, actor, "supplier.list_requested", "subcontract", sc.id)
    return e


def ensure_list_requests(db: Session) -> int:
    """Once per major-trade sub: when the job is first tracked or the sub first crosses the minimum."""
    n = 0
    for sc in db.scalars(select(Subcontract).join(Project).where(
            Project.is_tracked.is_(True), Project.track_suppliers.is_(True), Subcontract.deleted_in_source.is_(False),
            Subcontract.supplier_list_status.is_(None))).all():
        if is_major_trade(db, sc) and not active(db, sc):
            request_list(db, sc)
            n += 1
    return n


# ---------------------------------------------------------------- reading the reply

LIST_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["no_suppliers_over_minimum", "suppliers"],
    "properties": {
        "no_suppliers_over_minimum": {"type": "boolean",
                                      "description": "The sub says they have no suppliers over the minimum"},
        "suppliers": {"type": "array", "items": {
            "type": "object", "additionalProperties": False,
            "required": ["name", "supplies", "approx_amount", "contact_name", "email", "phone"],
            "properties": {k: {"type": ["string", "null"]} for k in
                           ("name", "supplies", "approx_amount", "contact_name", "email", "phone")}}},
    },
}

LIST_PROMPT = """A subcontractor replied to Buildout Pros' request for the suppliers and sub-tier
contractors on a job. List every supplier they named. Use null for anything they didn't give.
approx_amount: digits only. Don't include the subcontractor itself or Buildout Pros.
If they say they have no suppliers over {minimum}, set no_suppliers_over_minimum.

Their email:
{body}"""


def _blocks(db: Session, e: EmailMessage) -> list[dict]:
    """Attachments as Claude content blocks (PDFs, images, and spreadsheets as text)."""
    import base64
    import io

    from app.services import filestore

    out = []
    for a in e.attachments:
        data = filestore.read(db, a.path)
        if not data:
            continue
        if a.is_pdf:
            out.append({"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                       "data": base64.standard_b64encode(data).decode()}})
        elif a.is_image and (a.content_type or "") in ("image/png", "image/jpeg", "image/gif", "image/webp"):
            out.append({"type": "image", "source": {"type": "base64", "media_type": a.content_type,
                                                    "data": base64.standard_b64encode(data).decode()}})
        elif a.filename.lower().endswith((".xlsx", ".xlsm")):
            try:
                from openpyxl import load_workbook
                wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                rows = [", ".join("" if v is None else str(v) for v in r) for ws in wb.worksheets
                        for r in ws.iter_rows(values_only=True)]
                out.append({"type": "text", "text": f"Attached sheet {a.filename}:\n" + "\n".join(rows)[:20000]})
            except Exception:
                pass
        elif a.filename.lower().endswith((".csv", ".txt")):
            out.append({"type": "text", "text": f"Attached {a.filename}:\n" + data.decode("utf-8", "replace")[:20000]})
    return out


def read_list(db: Session, e: EmailMessage, client=None) -> dict | None:
    settings = get_settings()
    if client is None and not settings.anthropic_api_key:
        return None
    minimum = f"${Decimal(str(get_setting(db, 'supplier_min'))):,.0f}"
    content = _blocks(db, e) + [{"type": "text", "text": LIST_PROMPT.format(minimum=minimum,
                                                                           body=(e.body_text or "")[:20000])}]
    try:
        import anthropic

        client = client or anthropic.Anthropic(api_key=settings.anthropic_api_key)
        resp = client.beta.messages.create(
            model=settings.claude_model, max_tokens=4000, betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
            output_config={"effort": "low", "format": {"type": "json_schema", "schema": LIST_SCHEMA}},
            messages=[{"role": "user", "content": content}],
        )
        if resp.stop_reason == "refusal":
            return None
        return json.loads(next((b.text for b in resp.content if getattr(b, "type", "") == "text"), ""))
    except Exception as exc:
        log.warning("suppliers.read_list_failed", email_id=e.id, error=str(exc)[:200])
        return None


def handle_list_reply(db: Session, sc: Subcontract, e: EmailMessage, client=None) -> int:
    """A reply to the supplier-list request: each supplier becomes a yes/no question on Today."""
    result = read_list(db, e, client)
    if result is None:
        upsert_review(db, "supplier_list_unread", "email", e.id, sc.project.number,
                      f"{sc.vendor.name} replied about their supplier list. Open it and add the suppliers by hand.",
                      True)
        return 0
    if result.get("no_suppliers_over_minimum") and not result.get("suppliers"):
        sc.supplier_list_status = "none"
        return 0
    sc.supplier_list_status = "received"
    known = {normalize_vendor(s.name) for s in active(db, sc)}
    n = 0
    for row in result.get("suppliers") or []:
        name = (row.get("name") or "").strip()
        if not name or normalize_vendor(name) in known or normalize_vendor(name) == sc.vendor.normalized_name:
            continue
        known.add(normalize_vendor(name))
        db.add(SupplierCandidate(subcontract_id=sc.id, email_id=e.id, name=name[:255],
                                 supplies=(row.get("supplies") or None) and row["supplies"][:500],
                                 approx_amount=to_amount(row.get("approx_amount")),
                                 contact_name=row.get("contact_name"), email=row.get("email"),
                                 phone=(row.get("phone") or None) and row["phone"][:100]))
        n += 1
    db.flush()
    return n


def decide(db: Session, cand: SupplierCandidate, yes: bool, actor: str) -> Supplier | None:
    cand.status, cand.decided_by, cand.decided_at = ("added" if yes else "rejected"), actor, now()
    if not yes:
        return None
    return add(db, cand.subcontract, actor, cand.name, cand.supplies, cand.approx_amount, cand.contact_name,
               cand.email, cand.phone, source="list")


# ---------------------------------------------------------------- supplier waivers

def best_waiver(db: Session, s: Supplier) -> Waiver | None:
    """Same priority as a sub: unconditional final, else conditional final, else latest partial."""
    from datetime import date

    signed = db.scalars(select(Waiver).where(Waiver.supplier_id == s.id, Waiver.status.in_(SIGNED),
                                             Waiver.signed_pdf_path.is_not(None))).all()

    def recent(w: Waiver):
        return (w.through_date or date.min, w.received_at or now(), w.id)
    for t in (WaiverType.UNCONDITIONAL_FINAL, WaiverType.CONDITIONAL_FINAL):
        of_type = [w for w in signed if w.type == t.value]
        if of_type:
            return max(of_type, key=recent)
    partials = [w for w in signed if w.type in (WaiverType.CONDITIONAL_PARTIAL.value,
                                                WaiverType.UNCONDITIONAL_PARTIAL.value)]
    return max(partials, key=recent) if partials else None


def status_words(db: Session, s: Supplier) -> str:
    w = best_waiver(db, s)
    if w is None:
        return "No waiver yet"
    words = WaiverType(w.type).label
    if w.through_date and not WaiverType(w.type).is_final:
        words += f" through {w.through_date.month}/{w.through_date.day}"
    return words + (", signed" if w.status == WaiverStatus.ACCEPTED else ", signed (not reviewed yet)")


def match_supplier(db: Session, sc: Subcontract, company: str | None) -> Supplier | None:
    if not company:
        return None
    from rapidfuzz import fuzz

    norm = normalize_vendor(company)
    best, score = None, 0
    for s in active(db, sc):
        sn = normalize_vendor(s.name)
        if sn == norm:
            return s
        sc_ = fuzz.WRatio(norm, sn)
        if sc_ > score:
            best, score = s, sc_
    return best if score >= 90 else None


def file_supplier_waiver(db: Session, sub_waiver: Waiver, s: Supplier, att: EmailAttachment, result: dict,
                         actor: str) -> Waiver:
    """A supplier's signed waiver that came in with the sub's: filed under that supplier."""
    from pathlib import Path

    from app.mail.verify import compare_supplier
    from app.services import filestore
    from app.waivers import service as waivers

    found = result.get("found") or {}
    wtype = (found.get("waiver_type") or "").upper()
    if wtype not in {t.value for t in WaiverType}:
        wtype = sub_waiver.type
    w = db.scalar(select(Waiver).where(Waiver.payable_id == sub_waiver.payable_id, Waiver.type == wtype,
                                       Waiver.supplier_id == s.id))
    if w is None:
        w = Waiver(payable_id=sub_waiver.payable_id, subcontract_id=sub_waiver.subcontract_id, supplier_id=s.id,
                   type=wtype, state=sub_waiver.state, amount=to_amount(found.get("amount")),
                   through_date=sub_waiver.through_date, through_date_is_fallback=sub_waiver.through_date_is_fallback,
                   status=WaiverStatus.SENT.value, sent_at=sub_waiver.sent_at, requested_at=sub_waiver.requested_at,
                   reviewer_email=sub_waiver.reviewer_email)
        db.add(w)
        db.flush()
        waivers.log_event(db, w, actor, "created", note=f"supplier waiver from {s.name}, with {sub_waiver.subcontract.vendor.name}'s")
    data = filestore.read(db, att.path) or b""
    waivers.mark_received(db, w, actor, Path(att.filename).name, data)
    problems = compare_supplier(sub_waiver, found)
    w.verification = result | {"problems": problems, "status": "issues" if problems else "looks_good",
                               "email_id": att.email_id}
    recheck_finals(db, sub_waiver.subcontract)
    return w


# ---------------------------------------------------------------- finals and missing supplier waivers

def missing(db: Session, sc: Subcontract) -> list[Supplier]:
    return [s for s in active(db, sc) if best_waiver(db, s) is None]


def recheck_finals(db: Session, sc: Subcontract) -> None:
    """A sub's final lists the suppliers whose waivers are still missing (Quick question on Today)."""
    finals = db.scalars(select(Waiver).where(
        Waiver.subcontract_id == sc.id, Waiver.supplier_id.is_(None),
        Waiver.type.in_((WaiverType.CONDITIONAL_FINAL.value, WaiverType.UNCONDITIONAL_FINAL.value)),
        Waiver.status != WaiverStatus.VOID.value)).all()
    gone = missing(db, sc)
    for w in finals:
        if gone and w.status not in (WaiverStatus.ACCEPTED.value,):
            names = ", ".join(s.name for s in gone)
            upsert_review(db, "missing_supplier_waivers", "waiver", w.id, sc.project.number,
                          f"{sc.vendor.name}'s {WaiverType(w.type).label.lower()} is going out while supplier "
                          f"waivers are still missing from {names}.", True)
        else:
            upsert_review(db, "missing_supplier_waivers", "waiver", w.id, sc.project.number, "", False)


def set_tracking(db: Session, project: Project, on: bool, actor: str) -> int:
    """Turn suppliers on or off for a job. Turning off throws away unsent supplier-list emails
    and closes supplier questions; nothing already entered is deleted."""
    project.track_suppliers = on
    audit(db, actor, "project.suppliers_on" if on else "project.suppliers_off", "project", project.number)
    if on:
        for sc in project.subcontracts:
            recheck_finals(db, sc)
        return 0
    n = 0
    sc_ids = [sc.id for sc in project.subcontracts]
    for e in db.scalars(select(EmailMessage).where(
            EmailMessage.kind.in_(("supplier_list", "supplier_list_reminder")),
            EmailMessage.token.in_([f"SD-{i}" for i in sc_ids] or ["-"]),
            EmailMessage.status.in_(("draft", "queued", "failed")))):
        e.status = "discarded"
        n += 1
    for c in db.scalars(select(SupplierCandidate).where(SupplierCandidate.subcontract_id.in_(sc_ids or [-1]),
                                                        SupplierCandidate.status == "pending")):
        c.status, c.decided_by, c.decided_at = "rejected", f"{actor} (suppliers off for the job)", now()
    from app.db.models import ReviewItem
    for r in db.scalars(select(ReviewItem).where(ReviewItem.project_number == project.number, ReviewItem.status == "open",
                                                 ReviewItem.kind.in_(("missing_supplier_waivers", "supplier_list_unread",
                                                                      "supplier_waiver_unknown")))):
        r.status, r.resolved_by, r.resolved_at = "dismissed", actor, now()
        r.note = "Suppliers turned off for this job"
    return n
